import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def read_pdf(source) -> str:
    reader = PdfReader(source)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"

    return text


def read_docx(source) -> str:
    doc = Document(source)
    text = ""

    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text + "\n"

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def _extract(file_extension: str, source) -> str:
    if file_extension == ".pdf":
        return read_pdf(source)
    elif file_extension == ".docx":
        return read_docx(source)
    else:
        raise ValueError("Unsupported file format. Please provide a PDF or DOCX file.")


def read_resume(file_path: str) -> str:
    return _extract(Path(file_path).suffix.lower(), file_path)


def read_resume_bytes(file_name: str, data: bytes) -> str:
    """Extract text from an in-memory PDF/DOCX upload."""
    return _extract(Path(file_name).suffix.lower(), io.BytesIO(data))
